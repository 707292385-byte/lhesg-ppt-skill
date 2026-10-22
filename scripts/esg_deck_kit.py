# -*- coding: utf-8 -*-
"""
esg_deck_kit — lhesg 风格 PPT 构建工具库

核心思路：**基底自包含**
  本 skill 自带基底 assets/base.pptx（母版 + 版式 + 主题 + 封面/结尾页原型），
  构建时不依赖任何外部模板或既有 deck：以基底为底 → 清空两页原型 →
  用 layout2「标题和内容」新建内容页。视觉识别元素（标题蓝块/灰块、右上角logo、
  底部logo、三段式装饰条、标题字体与主题色）由母版与版式**原生继承**，无需逐页重绘。
  正文版式则完全自由设计（麦肯锡式结论标题 + 多样版式 + 图标）。

用法:
    from esg_deck_kit import EsgDeck, PRIMARY, BODY, F_HEAVY
    d = EsgDeck()                      # 基底内置，直接可用
    d.cover("主标题", presenter="姓名", date="2026年X月X日")  # 封面（按需；无副标题）
    d.toc(["第一章", "第二章", ...])    # 目录（按需）
    s = d.page("结论式标题：一句话说清本页观点")
    d.conclusion(s, "全页结论：紧贴标题下方，证据排在其下")   # 结论先行（按需）
    ... 用 d.text / d.rect / d.shape / d.icon / d.vline 排版证据 ...
    d.end("谢谢关注！", lines=["地址：…", "联系人：…"])   # 结尾页（按需）
    d.save("输出.pptx")

字号规范（唯一标准，lib 内强制）:
    正文类（观点/论据/表格内容/卡片正文/模块标题/关键数字/数据标签）不得小于 12pt，默认 14pt；
    注释类（口径说明、图注、单位、页脚来源行）用 9pt。set_font 对低于 9pt 直接报错。

命令行自检:
    python esg_deck_kit.py --check      # 打印基底自检信息
"""
import os
import sys
import tempfile
import re
import copy
import shutil
import zipfile

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LEGEND_POSITION,
                             XL_LABEL_POSITION, XL_TICK_MARK,
                             XL_MARKER_STYLE)
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn, nsdecls


def _xml_escape(s):
    return (s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_ROOT = os.path.dirname(HERE)
ICON_DIR = os.path.join(SKILL_ROOT, "assets", "icons")

# ── lhesg 主题色（基底 layout2 / 母版定义值）──
PRIMARY = RGBColor.from_string("044AC0")   # 标题主蓝（layout2 title defRPr）
NAVY    = RGBColor.from_string("123685")   # 深蓝，表头/重块
DEEP    = RGBColor.from_string("0F4A9E")   # 次级深蓝
BRIGHT  = RGBColor.from_string("376FFF")   # 亮蓝，竖线/强调
BODY    = RGBColor.from_string("333333")   # 正文
GRAY    = RGBColor.from_string("595959")   # 次级说明
FOOT    = RGBColor.from_string("7F7F7F")   # 来源行
LIGHT   = RGBColor.from_string("EEF3FC")   # 卡片浅蓝底
BORDER  = RGBColor.from_string("D6E2F5")   # 描边
MUTED   = RGBColor.from_string("8FAADC")   # 弱化蓝（对比中的"少数/次要"方）
ORANGE  = RGBColor.from_string("C55A11")   # 强调橙（警示/负面，克制使用）
WHITE   = RGBColor.from_string("FFFFFF")
ROWALT  = RGBColor.from_string("F2F6FC")   # 表格隔行

# 图表系列色序：单系列用 PRIMARY；多系列按此顺序取，橙色不进默认序列（留给负面强调）
CHART_SERIES = (PRIMARY, BRIGHT, MUTED, DEEP)

F_HEAVY, F_MED, F_LIGHT = "思源黑体 Heavy", "思源黑体 Medium", "思源黑体 Light"

# ── 字号规范（唯一标准：正文不小于 12pt；注释与来源可用 9pt）──
# 正文类（观点/论据/表格内容/卡片正文/模块标题/关键数字/数据标签）：≥ 12pt
# 注释类（口径说明、图注、单位、页脚来源行）：9pt
S_TITLE = 24     # 页标题（母版 layout2 固定，勿改）
S_H2    = 15     # 卡片 / 模块标题（Medium Bold）
S_BODY  = 14     # 正文默认字号 ★
S_DENSE = 12     # 正文下限——仅信息密集页使用
S_AUX   = 9      # 注释 / 口径说明 / 单位 / 页脚来源行
S_COVER = 28     # 封面标题（基底封面原值，勿改）
S_END   = 44     # 结尾页主标题（基底结尾页原值，勿改）
S_NUM   = 30     # 关键数字
S_NUM_L = 34     # 数据卡主数字

# 合法字号档位（font_lint.py 依此校验；不在此集合内即为非规范档位）
SIZE_LADDER = (S_AUX, S_DENSE, S_BODY, S_H2, S_TITLE, S_COVER, S_NUM, S_NUM_L, S_END)

# ── 画布与版式安全区（16:9, 13.333in × 7.5in）──
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN_X = 0.67            # 左边界（与 layout2 标题占位符对齐）
CONTENT_W = 11.99          # 正文可用宽度
CONTENT_TOP = 1.42         # 正文起始 y（标题行下方）
CONTENT_BOTTOM = 6.80      # 正文结束 y（来源行上沿）
SRC_Y = 6.88               # 底部来源行 y
TITLE_LEFT, TITLE_TOP = 0.67, 0.44
TITLE_W, TITLE_H = 11.99, 0.55

# ── 页脚来源行默认值：每套 deck 统一改这一处 ──
SRC_LINE = "资料来源：（逐条可回溯到原件，禁止臆造）"

BASE_PPTX = os.path.join(SKILL_ROOT, "assets", "base.pptx")
BASE_ENV = "ESG_BASE"   # 自定义基底时用环境变量覆盖，正常无需设置

TRACK = RGBColor.from_string("E3EAF6")     # 比例条底槽
BLUE_BG = RGBColor.from_string("1C589E")   # 结尾页深蓝底（基底 slide28 实测）
BLUE_PANEL = RGBColor.from_string("1F5FAD")  # 结尾页装饰面板（略浅一档）

# ── 分节清除：presentation.xml 的 p14:sectionLst 独立于页码，删页不会删它；
#    其 slideId 指向已删页时，打开文件分节视图会出现残留节名，可能触发修复提示。
#    基底已预清，建 deck 时再清一次是保险（外部 deck 用 strip_sections.py）。──
SECTION_EXT_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"
SEC_PATH = ".//{%s}section" % P14_NS


_TINT_CACHE_DIR = os.path.join(tempfile.gettempdir(), "lhesg_ppt_tint")


def _tinted_icon(path, tint):
    """把图标重染为指定颜色（保留 alpha），返回缓存文件路径。"""
    from PIL import Image
    key = "%s_%s.png" % (os.path.splitext(os.path.basename(path))[0],
                         str(tint).replace(" ", ""))
    cache = os.path.join(_TINT_CACHE_DIR, key)
    if not os.path.exists(cache):
        os.makedirs(_TINT_CACHE_DIR, exist_ok=True)
        im = Image.open(path).convert("RGBA")
        px = im.getdata()
        out = [(tint[0], tint[1], tint[2], a) for (_, _, _, a) in px]
        im.putdata(out)
        im.save(cache)
    return cache


def strip_sections(xml_root):
    """移除 presentation.xml 里的分节定义，返回 (移除的 ext 数, 移除的节数)。"""
    n_ext = n_sec = 0
    for ext_lst in xml_root.findall(qn("p:extLst")):
        for ext in list(ext_lst.findall(qn("p:ext"))):
            secs = ext.findall(SEC_PATH)
            if ext.get("uri") == SECTION_EXT_URI or secs:
                n_sec += len(secs)
                ext_lst.remove(ext)
                n_ext += 1
        if len(ext_lst) == 0:
            ext_lst.getparent().remove(ext_lst)
    return n_ext, n_sec


def find_base(explicit=None):
    """基底定位：参数 → 环境变量 ESG_BASE → skill 自带 assets/base.pptx。
    基底随 skill 分发，正常情况下不依赖任何外部文件。"""
    if explicit and os.path.isfile(explicit):
        return os.path.abspath(explicit)
    env = os.environ.get(BASE_ENV)
    if env and os.path.isfile(env):
        return os.path.abspath(env)
    if os.path.isfile(BASE_PPTX):
        return BASE_PPTX
    raise FileNotFoundError(
        "缺少基底 assets/base.pptx（应随 skill 分发，勿删除）。"
        "如用自定义基底，设环境变量 ESG_BASE 指向其路径。"
    )


LEGACY_ICONS = ["image21.png", "image23.png", "image25.png", "image27.png",
                "image29.png", "image31.png", "image33.png"]


def ensure_icons(base_path):
    """返回 assets/icons 下全部图标文件名（100 个：7 个模板原生 + 93 个同风格
    Material 实心图标 icon_*.png，均 800×800）。7 个模板原生图标缺失时从基底
    media 兜底提取。"""
    os.makedirs(ICON_DIR, exist_ok=True)
    missing = [n for n in LEGACY_ICONS if not os.path.exists(os.path.join(ICON_DIR, n))]
    if missing:
        with zipfile.ZipFile(base_path) as z:
            for n in missing:
                try:
                    with open(os.path.join(ICON_DIR, n), "wb") as f:
                        f.write(z.read("ppt/media/" + n))
                except KeyError:
                    pass
    names = sorted(f for f in os.listdir(ICON_DIR) if f.lower().endswith(".png"))
    return names


# ── 品牌图：封面 / 结尾页用，随 skill 分发，缺失时从基底提取（兜底） ──
BRAND_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "brand")
BRAND_SRC = {  # skill 资产名 -> 基底 media 来源（兜底提取用，资产本身随 skill 分发）
    "cover_bg.png": "ppt/media/image3.png",     # 封面地球背景（含 lh 水印）
    "banner_arrow.png": "ppt/media/image4.png",  # 封面标题箭头条
    "logo_white.png": "ppt/media/image5.png",    # 白色全称 logo（封面母版层）
    "logo_blue.png": "ppt/media/image47.png",    # 蓝色 logo（结尾页）
    "qr_public.jpeg": "ppt/media/image48.jpeg",  # 公众号二维码（结尾页）
}


def ensure_brand(base_path):
    os.makedirs(BRAND_DIR, exist_ok=True)
    missing = [n for n in BRAND_SRC if not os.path.exists(os.path.join(BRAND_DIR, n))]
    if missing:
        with zipfile.ZipFile(base_path) as z:
            for n in missing:
                try:
                    with open(os.path.join(BRAND_DIR, n), "wb") as f:
                        f.write(z.read(BRAND_SRC[n]))
                except KeyError:
                    pass
    return {os.path.splitext(n)[0]: os.path.join(BRAND_DIR, n) for n in BRAND_SRC
            if os.path.exists(os.path.join(BRAND_DIR, n))}


# ══════════════════════════════ 绘图原语 ══════════════════════════════

def set_font(run, size=14, color=BODY, name=F_LIGHT, bold=False):
    """中文字体必须同时落 latin / ea / cs 三个 typeface，仅设 font.name 只写 latin。
    color=None 表示不设色（沿用主题默认，仅封面等沿用基底原值时使用）。"""
    if size < S_AUX:
        raise ValueError(
            f"字号 {size}pt 低于规范下限：正文不小于 {S_DENSE}pt，注释/来源不小于 {S_AUX}pt")
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.name = name
    if color is not None:
        f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", name)


def _norm_runs(p):
    """允许 (text,size,color,font) 四元组，自动补 bold=False。"""
    return [tuple(list(r) + [False])[:5] if len(r) == 4 else r for r in p]


class EsgDeck:
    """lhesg 风格 deck 构建器。所有几何单位均为英寸。"""

    def __init__(self, base=None, source_line=None):
        self.base = find_base(base)
        self.icons = ensure_icons(self.base)
        self.brand = ensure_brand(self.base)
        self.source_line = source_line or SRC_LINE
        self.prs = Presentation(self.base)
        self._clear_slides()
        self.layout = self._pick_layout("slideLayout2.xml")
        self.layout_title = self._pick_layout("slideLayout1.xml")
        self.slides = []

    # ── 初始化 ──
    def _clear_slides(self):
        """清空基底两页原型（封面/结尾页克隆数据直接读 self.base 文件，不受影响）。
        必须 drop_rel 再 remove，否则残留关系导致文件损坏。
        同时清掉残留的分节定义（基底已预清，此处是保险）。"""
        sldIdLst = self.prs.slides._sldIdLst
        for sldId in list(sldIdLst):
            rId = sldId.get(qn("r:id"))
            self.prs.part.drop_rel(rId)
            sldIdLst.remove(sldId)
        strip_sections(self.prs.part._element)

    def _pick_layout(self, partname_suffix):
        for m in self.prs.slide_masters:
            for lay in m.slide_layouts:
                if lay.part.partname.endswith(partname_suffix):
                    return lay
        raise KeyError(f"基底缺少版式 {partname_suffix}")

    # ── 基底原页克隆（封面/结尾页用，保证与内置原型完全一致）──
    @staticmethod
    def _norm_txt(t):
        return re.sub(r"\s+", "", t or "")

    def _clone_template_slide(self, slide_xml, texts=None):
        """克隆基底某一页：版式、p:bg 背景、全部形状（图片/组合/freeform）、
        字号颜色原样保留，仅按 texts 替换文字（键=去空白后的旧文本，值=新文本，
        新文本含 \n 时拆多段，沿用原 run 的字体格式）。"""
        with zipfile.ZipFile(self.base) as z:
            sx = z.read("ppt/slides/" + slide_xml).decode("utf-8")
            rels = z.read("ppt/slides/_rels/" + slide_xml + ".rels").decode("utf-8")
            media = []
            for rid, tgt in re.findall(r'Id="(rId\d+)"[^>]*Target="\.\./(media/[^"]+)"', rels):
                media.append((rid, tgt, z.read("ppt/" + tgt)))
        lay = self._pick_layout(
            re.search(r'Target="\.\./(slideLayouts/[^"]+)"', rels).group(1).split("/")[-1])
        s = self.prs.slides.add_slide(lay)
        # 图片关系重映射（两段式替换，防新旧 rId 撞号）
        if media:
            for rid, _, _ in media:
                sx = sx.replace(f'r:embed="{rid}"', f'r:embed="@@{rid}@@"')
                sx = sx.replace(f'r:link="{rid}"', f'r:link="@@{rid}@@"')
            tmp_files = []
            for rid, tgt, data in media:
                fd, path = tempfile.mkstemp(suffix=os.path.splitext(tgt)[1])
                with os.fdopen(fd, "wb") as f:
                    f.write(data)
                tmp_files.append(path)
                _, new_rid = s.part.get_or_add_image_part(path)
                sx = sx.replace(f'"@@{rid}@@"', f'"{new_rid}"')
            for path in tmp_files:
                os.unlink(path)
        # 用基底页的 cSld 内容（bg + spTree）替换新页的
        tpl_root = parse_xml(sx.encode("utf-8"))
        tpl_cSld = tpl_root.find(qn("p:cSld"))
        my_cSld = s._element.find(qn("p:cSld"))
        my_tree = my_cSld.find(qn("p:spTree"))
        for child in list(my_tree):
            if child.tag in (qn("p:sp"), qn("p:pic"), qn("p:graphicFrame"),
                             qn("p:grpSp"), qn("p:cxnSp")):
                my_tree.remove(child)
        tpl_bg = tpl_cSld.find(qn("p:bg"))
        if tpl_bg is not None:
            old_bg = my_cSld.find(qn("p:bg"))
            if old_bg is not None:
                my_cSld.remove(old_bg)
            my_cSld.insert(0, copy.deepcopy(tpl_bg))
        tpl_tree = tpl_cSld.find(qn("p:spTree"))
        for child in tpl_tree:
            if child.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
                continue
            my_tree.append(copy.deepcopy(child))
        # 文字替换
        if texts:
            def visit(shapes):
                for sh in shapes:
                    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                        visit(sh.shapes)
                        continue
                    if getattr(sh, "has_text_frame", False):
                        key = self._norm_txt(sh.text_frame.text)
                        if key in texts:
                            self._set_shape_text(sh, texts[key])
            visit(s.shapes)
        self.slides.append(s)
        return s

    def _fit_text_width(self, slide, text, width_in, no_wrap=True):
        """把承载指定文字的文本框放宽到 width_in（页内英寸），并关闭自动换行。

        用于封面「汇报人」、结尾页「联系人」这类窄框：原宽仅 1.13in，
        5 个汉字加左右边距就会折成两行。组内文本框按 group 的
        chExt→ext 比例换算成子坐标宽度。no_wrap 兜底：即便值再长也不折行。"""
        key = self._norm_txt(text)

        def walk(shapes, scale):
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    xfrm = sh._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
                    ext, ch = xfrm.find(qn("a:ext")), xfrm.find(qn("a:chExt"))
                    walk(sh.shapes, scale * int(ext.get("cx")) / int(ch.get("cx")))
                    continue
                if not getattr(sh, "has_text_frame", False):
                    continue
                if self._norm_txt(sh.text_frame.text) != key:
                    continue
                sh.width = Emu(int(round(width_in * 914400 / scale)))
                if no_wrap:
                    bp = sh.text_frame._txBody.find(qn("a:bodyPr"))
                    if bp is not None:
                        bp.set("wrap", "none")
                return True
            return False

        return walk(slide.shapes, 1.0)

    @staticmethod
    def _set_shape_text(sh, new_text):
        """替换形状文字：沿用原第一个 run 的字体格式；\n 拆多段。"""
        tf = sh.text_frame
        proto = None
        for para in tf.paragraphs:
            if para.runs:
                proto = copy.deepcopy(para.runs[0]._r.get_or_add_rPr())
                break
        for para in list(tf.paragraphs)[1:]:
            para._p.getparent().remove(para._p)
        p0 = tf.paragraphs[0]
        for r in list(p0.runs):
            r._r.getparent().remove(r._r)
        last = p0._p
        for i, ln in enumerate(new_text.split("\n")):
            if i == 0:
                para_el = p0._p
            else:
                para_el = copy.deepcopy(p0._p)
                for r_el in para_el.findall(qn("a:r")):
                    para_el.remove(r_el)
                last.addnext(para_el)
            last = para_el
            run = parse_xml("<a:r %s><a:t>%s</a:t></a:r>" % (nsdecls("a"), _xml_escape(ln)))
            if proto is not None:
                run.insert(0, copy.deepcopy(proto))
            para_el.append(run)

    # ── 建页 ──
    def page(self, title_text, source=False):
        """新建一页：保留 layout2 的标题与页码占位符，删除正文占位符，正文区自由排版。

        source 默认 False——**只有引用外部材料才加注释**；本地材料（客户官网、
        内部台账、自己整理的分类表、过程文件）不加来源行。需要加时用 d.note()。"""
        s = self.prs.slides.add_slide(self.layout)
        for ph in list(s.placeholders):
            t = ph.placeholder_format.type
            if t == PP_PLACEHOLDER.TITLE:
                continue
            if t == PP_PLACEHOLDER.SLIDE_NUMBER:
                continue
            ph._element.getparent().remove(ph._element)
        title = s.shapes.title
        title.left, title.top = Inches(TITLE_LEFT), Inches(TITLE_TOP)
        title.width, title.height = Inches(TITLE_W), Inches(TITLE_H)
        tf = title.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = tf.paragraphs[0].add_run()
        r.text = title_text
        set_font(r, 24, PRIMARY, F_HEAVY)
        if source:
            self.text(s, MARGIN_X, SRC_Y, CONTENT_W, 0.25,
                      [[(self.source_line, S_AUX, FOOT, F_LIGHT)]])
        self.slides.append(s)
        return s

    def conclusion(self, slide, text, y=CONTENT_TOP, h=0.88,
                   fill=None, accent=PRIMARY, size=S_BODY):
        """全页结论：**结论先行**——紧贴标题之下、证据之上放一行文字。

        不加底色块、不加竖线、不加框——结论就是一行排版的文字，
        靠字重和颜色区分，不靠色块。
        页面纵向顺序固定为：**标题 → 全页结论 → 证据 → 来源行**。

        text 三种写法：
          - 字符串：整句按"主蓝加粗"排版；
          - 单个段落（run 列表）：[(文本, 字号, 颜色, 字体[, 加粗]), ...]，
            常用做法是"前半句判断加粗主蓝 + 后半句说明常规"：
            [(判断句, S_BODY, PRIMARY, F_MED, True), (说明, S_BODY, BODY, F_LIGHT)]
          - 多段落：[[run, ...], [run, ...]]

        返回结论区底边 y；证据区从 返回的 y + 0.24 起排。"""
        if isinstance(text, str):
            paras = [[(text, size, accent, F_MED, True)]]
        elif text and isinstance(text[0], tuple):
            paras = [text]
        else:
            paras = text
        self.text(slide, MARGIN_X, y, CONTENT_W, h, paras,
                  line_spacing=1.2)
        return y + h

    def note(self, slide, text, y=SRC_Y, w=CONTENT_W):
        """页脚注释（9pt 灰字）。**只用于引用外部材料**：外部报告/标准/研究的
        核心结论、数据、图表，需写清来源名称、年份与条款或页码。

        本地来源材料（客户官网、内部台账、自制分类表、过程文件）不加注释；
        没有外部引用就不调用本方法。"""
        return self.text(slide, MARGIN_X, y, w, 0.26,
                         [[(text, S_AUX, FOOT, F_LIGHT)]])

    # ── 特殊页：封面 / 目录 / 结尾 ──
    def _strip_placeholders(self, slide, keep_title=True):
        """只保留标题与页码占位符，其余（正文/日期/页脚）删掉，避免空占位符留框。"""
        for ph in list(slide.placeholders):
            t = ph.placeholder_format.type
            if t == PP_PLACEHOLDER.SLIDE_NUMBER:
                continue
            if keep_title and t == PP_PLACEHOLDER.TITLE:
                continue
            ph._element.getparent().remove(ph._element)

    def _set_title(self, slide, text, x, y, w, h, size=24,
                   color=PRIMARY, font=None, bold=False, align=PP_ALIGN.LEFT,
                   line_spacing=1.15):
        """标题写入占位符。text 含 "\\n" 时拆成多段（如封面两行标题）。"""
        title = slide.shapes.title
        title.left, title.top = Inches(x), Inches(y)
        title.width, title.height = Inches(w), Inches(h)
        tf = title.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, seg in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line_spacing
            r = p.add_run()
            r.text = seg
            set_font(r, size, color, font or F_HEAVY, bold)
        return title

    def _set_bg_picture(self, slide, img_path):
        """把整页背景设成图片（写 p:bg blipFill，与基底封面做法一致）。
        用背景而非全幅图片形状，母版的 logo 与装饰条才会保留在上层。"""
        _, rId = slide.part.get_or_add_image_part(img_path)
        bg = parse_xml(
            '<p:bg %s><p:bgPr><a:blipFill rotWithShape="1">'
            '<a:blip r:embed="%s"/><a:srcRect/><a:stretch><a:fillRect/>'
            '</a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>'
            % (nsdecls("p", "a", "r"), rId))
        cSld = slide._element.find(qn("p:cSld"))
        cSld.insert(0, bg)

    def cover(self, title, presenter=None, date=None):
        """封面页：整页克隆基底 slide1 原型（地球背景/箭头条/白色 logo/汇报人框原样），
        仅替换文字；箭头条加高到 1.89in、标题框随之加高并垂直居中，容纳两行
        28pt 标题（title 含 \n 即两行，超宽自动折行）。字号颜色沿用基底原值。
        封面不含副标题——基底原页没有该元素，需要补充信息就写进标题的次行。"""
        texts = {"此外添加标题": title}
        if presenter:
            texts["XXX"] = presenter
        if date:
            texts["20XX年XX月XX日"] = date
        s = self._clone_template_slide("slide1.xml", texts)
        BANNER_H = 1.89
        title_head = self._norm_txt(title.split("\n")[0])
        for sh in s.shapes:
            try:
                L, T, W, H = (round(v / 914400, 2) for v in
                              (sh.left, sh.top, sh.width, sh.height))
            except Exception:
                continue
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and \
               abs(W - 7.54) < 0.05 and abs(H - 1.23) < 0.05:
                # 箭头条加高（中心 2.495 不变），两行 28pt 稳稳放进条内
                sh.top = Inches(2.495 - BANNER_H / 2)
                sh.height = Inches(BANNER_H)
            elif getattr(sh, "has_text_frame", False) and \
                    title_head and title_head in self._norm_txt(sh.text_frame.text):
                sh.top = Inches(1.70)
                sh.height = Inches(1.60)
                bp = sh.text_frame._txBody.find(qn("a:bodyPr"))
                bp.set("anchor", "ctr")
        # 汇报人值框原宽仅 1.13in，5 字即折行 → 放宽到与日期行等宽 2.24in
        if presenter:
            self._fit_text_width(s, presenter, 2.24)
        return s

    def toc(self, items, title="目录"):
        """目录页：沿用内容页标题样式（蓝块+灰块），正文区为编号卡片列表。
        items 元素可为 "篇章名" 或 ("篇章名", "右侧注释")；卡片浅蓝底与白底交替。
        特殊页不设副标题槽——标题即结论，页面文本只放内容。"""
        s = self.page(title, source=False)
        n = max(len(items), 1)
        top, bottom = 1.72, 6.72
        step = min(0.98, (bottom - top) / n)
        h = min(0.74, step - 0.16)
        for i, it in enumerate(items):
            name, note = it if isinstance(it, (tuple, list)) else (it, None)
            y = top + i * step
            if i % 2 == 0:
                self.rect(s, MARGIN_X, y, CONTENT_W, h, fill=LIGHT,
                          rounded=True, radius=0.12)
            chip = self.rect(s, MARGIN_X + 0.26, y + (h - 0.50) / 2, 0.50, 0.50,
                             fill=PRIMARY, rounded=True, radius=0.20)
            self.shape_text(chip, [[("%02d" % (i + 1), S_H2, WHITE, F_MED, True)]])
            self.text(s, MARGIN_X + 0.98, y, 7.20, h,
                      [[(name, S_H2, BODY, F_MED, True)]], anchor=MSO_ANCHOR.MIDDLE)
            if note:
                self.text(s, 8.60, y, CONTENT_W - 7.93, h,
                          [[(note, S_AUX, GRAY, F_LIGHT)]],
                          anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
        return s

    def end(self, title="谢谢关注！", lines=None, contact=None):
        """结尾页：整页克隆基底 slide28 原型（深蓝底/装饰弧面/公众号二维码/底部白区
        蓝色 logo/字号颜色全部原样），仅替换文字。
        lines: 地址等信息行列表（None=保留基底原文）；
        contact: [联系人, 联系电话] 两个值（None=保留基底占位）。"""
        texts = {"谢谢关注！": title}
        if contact:
            if contact[0]:
                texts["XXX"] = contact[0]
            if len(contact) > 1 and contact[1]:
                texts["88888888888"] = contact[1]
        s = self._clone_template_slide("slide28.xml", texts)
        # 联系人值框原宽仅 1.13in（同级电话框 2.235in）→ 放到组右边界
        if contact and contact[0]:
            self._fit_text_width(s, contact[0], 2.23)
        if lines:
            for sh in s.shapes:
                if getattr(sh, "has_text_frame", False) and \
                        self._norm_txt(sh.text_frame.text).startswith("地址"):
                    self._set_shape_text(sh, "\n".join(lines))
        return s

    # ── 文字 ──
    def text(self, slide, x, y, w, h, paras, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.15, space_after=4):
        """paras: 段落列表；每段是 [(文本, 字号, 颜色, 字体[, 加粗]), ...]"""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        first = True
        for p in paras:
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            para.alignment = align
            para.line_spacing = line_spacing
            para.space_after = Pt(space_after)
            for (txt, sz, clr, fn, bd) in _norm_runs(p):
                r = para.add_run()
                r.text = txt
                set_font(r, sz, clr, fn, bd)
        return tb

    # ── 形状 ──
    def rect(self, slide, x, y, w, h, fill=None, line=None,
             rounded=True, radius=0.05, line_w=0.75):
        """色块默认**轻微圆角**（radius=0.05，约为短边的 5%）。要直角显式传 rounded=False。"""
        shp = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h))
        if rounded:
            try:
                shp.adjustments[0] = radius
            except Exception:
                pass
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = fill
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line
            shp.line.width = Pt(line_w)
        shp.shadow.inherit = False
        return shp

    def shape(self, slide, kind, x, y, w, h, fill, line=None, radius=None):
        shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius is not None:
            try:
                shp.adjustments[0] = radius
            except Exception:
                pass
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line
            shp.line.width = Pt(0.75)
        shp.shadow.inherit = False
        return shp

    def shape_text(self, shp, paras, anchor=MSO_ANCHOR.MIDDLE,
                   align=PP_ALIGN.CENTER, line_spacing=1.1):
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.08)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        first = True
        for p in paras:
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            para.alignment = align
            para.line_spacing = line_spacing
            for (txt, sz, clr, fn, bd) in _norm_runs(p):
                r = para.add_run()
                r.text = txt
                set_font(r, sz, clr, fn, bd)

    def vline(self, slide, x, y, h, color=BRIGHT, w=0.045):
        """卡片左侧竖强调线（本风格惯用）。"""
        return self.rect(slide, x, y, w, h, fill=color)

    def icon(self, slide, icon, x, y, size, circle_fill=None, tint=None):
        """插入图标（assets/icons/ 下的文件名，如 "icon_eco.png"）。

        - circle_fill=None：图标直接放置（适合蓝系图标）。
        - circle_fill=颜色：先画圆形底，图标内缩 22% 叠加（白图标垫蓝圆底）。
        - tint=颜色（如 WHITE）：把非透明像素整体重染为该色（保留 alpha），
          蓝色 icon_*.png 也可垫蓝圆底使用；重染结果缓存在临时目录。"""
        path = os.path.join(ICON_DIR, icon)
        if tint is not None:
            path = _tinted_icon(path, tint)
        if circle_fill is not None:
            self.shape(slide, MSO_SHAPE.OVAL, x, y, size, size, circle_fill)
            pad = size * 0.22
            slide.shapes.add_picture(path, Inches(x + pad), Inches(y + pad),
                                     Inches(size - 2 * pad), Inches(size - 2 * pad))
        else:
            slide.shapes.add_picture(path, Inches(x), Inches(y),
                                     Inches(size), Inches(size))

    def table(self, slide, x, y, w, h, data, col_w=None,
              header_fill=NAVY, row_h=None, header_h=0.62,
              center_cols=(), body_size=S_DENSE, header_size=S_BODY):
        """表格：表头深蓝白字；正文必须显式 LEFT（默认表格样式会把 CJK 右对齐）。"""
        rows, cols = len(data), len(data[0])
        tb = slide.shapes.add_table(rows, cols, Inches(x), Inches(y),
                                    Inches(w), Inches(h)).table
        if col_w:
            for i, cw in enumerate(col_w):
                tb.columns[i].width = Inches(cw)
        if row_h:
            for ri in range(1, rows):
                tb.rows[ri].height = Inches(row_h)
        tb.rows[0].height = Inches(header_h)
        for ri in range(rows):
            for ci in range(cols):
                cell = tb.cell(ri, ci)
                cell.margin_left = cell.margin_right = Inches(0.12)
                cell.margin_top = cell.margin_bottom = Inches(0.06)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                for li, ln in enumerate(str(data[ri][ci]).split("\n")):
                    para = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                    para.alignment = PP_ALIGN.CENTER if (ri == 0 or ci in center_cols) else PP_ALIGN.LEFT
                    r = para.add_run()
                    r.text = ln
                    if ri == 0:
                        set_font(r, header_size, WHITE, F_MED, True)
                    elif ci == 0:
                        set_font(r, body_size + 2 if li == 0 else body_size,
                                 PRIMARY if li == 0 else GRAY,
                                 F_MED if li == 0 else F_LIGHT, li == 0)
                    else:
                        set_font(r, body_size, BODY, F_LIGHT)
                if ri == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = header_fill
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = WHITE if ri % 2 else ROWALT
        return tb

    # ── 图表 ──
    _CHART_KINDS = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "area": XL_CHART_TYPE.AREA,
        "pie": XL_CHART_TYPE.PIE,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
    }

    def chart(self, slide, kind, categories, series, x, y, w, h,
              colors=None, highlight=None, highlight_color=BRIGHT,
              other_color=MUTED, number_format="0",
              value_size=S_DENSE, cat_size=S_DENSE,
              gap_width=55, show_values=True, legend=False,
              val_min=0, val_max=None, reverse_cats=None,
              line_w=2.25, label_pos=None, show_cat_axis=True,
              show_val_axis=False):
        """原生图表：横向条形 / 柱 / 折线 / 面积 / 饼。

        kind         'bar'（横向条形，类目多时首选）/ 'column' / 'line' /
                     'area' / 'pie' / 'doughnut'
        categories   类目列表（已排序好再传入，函数不排序）
        series       [(系列名, 数值列表), ...]；单系列可直接传数值列表
        highlight    需要强调的点下标（int 或 list）。给定时其余点转 MUTED，
                     只留一处亮色——这是数据页最有效的"反模板"手段：
                     读者先看到那一条，再回看整体。
        reverse_cats None 时横向条形自动反转（降序数据的第一个排最上），
                     其余图表保持自然顺序。

        风格内置：无网格线、无图表边框、单系列无图例、数值直标、零基线、
        类目轴只留一条浅色轴线。不做 3D、不用阴影、不用彩虹色。
        """
        if kind not in self._CHART_KINDS:
            raise ValueError(f"unsupported kind: {kind}")
        if series and isinstance(series[0], (int, float)):
            series = [("", list(series))]
        cd = CategoryChartData()
        cd.categories = list(categories)
        for name, values in series:
            cd.add_series(name or "S1", list(values))

        gf = slide.shapes.add_chart(self._CHART_KINDS[kind], Inches(x), Inches(y),
                                    Inches(w), Inches(h), cd)
        ch = gf.chart
        ch.has_title = False
        ch.has_legend = bool(legend)
        if legend:
            ch.legend.position = XL_LEGEND_POSITION.RIGHT
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(S_AUX)
            ch.legend.font.name = F_LIGHT
            ch.legend.font.color.rgb = GRAY

        # 图表区：无填充无边框，融进页面而非"贴一张图"
        # python-pptx 未暴露 chartSpace 的 spPr，直接写 XML
        cs = ch._chartSpace
        spPr = parse_xml('<c:spPr %s><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
                         % nsdecls("c", "a"))
        anchor_el = cs.find(qn("c:txPr"))
        if anchor_el is None:
            anchor_el = cs.find(qn("c:externalData"))
        if anchor_el is not None:
            anchor_el.addprevious(spPr)
        else:
            cs.append(spPr)
        ch.font.size = Pt(cat_size)
        ch.font.name = F_LIGHT
        ch.font.color.rgb = GRAY

        plot = ch.plots[0]
        is_pie = kind in ("pie", "doughnut")
        if not is_pie:
            plot.gap_width = gap_width
            try:
                plot.overlap = -10 if len(series) > 1 else 0
            except Exception:
                pass
        try:
            plot.vary_by_categories = False
        except Exception:
            pass

        # 类目轴 / 数值轴：网格线全关，只留一条浅轴线作基线
        cax = ch.category_axis
        cax.has_major_gridlines = False
        cax.has_minor_gridlines = False
        cax.has_title = False
        try:
            cax.major_tick_mark = XL_TICK_MARK.NONE
            cax.minor_tick_mark = XL_TICK_MARK.NONE
        except Exception:
            pass
        try:
            cax.format.line.color.rgb = BORDER
            cax.format.line.width = Pt(0.75)
        except Exception:
            pass
        try:
            cax.tick_labels.font.size = Pt(cat_size)
            cax.tick_labels.font.name = F_LIGHT
            cax.tick_labels.font.color.rgb = GRAY
        except Exception:
            pass
        if not show_cat_axis:
            cax.visible = False
        elif reverse_cats is None:
            reverse_cats = (kind == "bar")
        if reverse_cats:
            # 反转类目顺序：XML 层设置，兼容 python-pptx 未暴露该属性
            scaling = cax._element.find(qn("c:scaling"))
            orient = scaling.find(qn("c:orientation"))
            if orient is not None:
                orient.set("val", "maxMin")

        vax = ch.value_axis
        vax.has_major_gridlines = False
        vax.has_minor_gridlines = False
        vax.has_title = False
        if not is_pie:
            try:
                vax.minimum_scale = val_min
                if val_max is not None:
                    vax.maximum_scale = val_max
            except Exception:
                pass
        vax.visible = bool(show_val_axis)
        if show_val_axis:
            vax.tick_labels.font.size = Pt(S_AUX)

        # 数据标签：直接标数值，不靠图例和刻度换算
        if show_values:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.number_format = number_format
            dl.number_format_is_linked = False
            dl.font.size = Pt(value_size)
            dl.font.name = F_MED
            dl.font.color.rgb = GRAY
            if is_pie:
                dl.position = XL_LABEL_POSITION.CENTER
            elif label_pos is not None:
                dl.position = label_pos
            elif kind in ("bar", "column"):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            elif kind == "line":
                dl.position = XL_LABEL_POSITION.ABOVE

        # 配色
        palette = list(colors) if colors else list(CHART_SERIES)
        idxs = set()
        if highlight is not None:
            idxs = {highlight} if isinstance(highlight, int) else set(highlight)
        for si, ser in enumerate(ch.series):
            base = palette[si % len(palette)]
            if kind == "line":
                ser.format.line.color.rgb = base
                ser.format.line.width = Pt(line_w)
                try:
                    ser.smooth = False
                except Exception:
                    pass
                try:
                    ser.marker.style = XL_MARKER_STYLE.CIRCLE
                    ser.marker.size = 7
                    ser.marker.format.fill.solid()
                    ser.marker.format.fill.fore_color.rgb = base
                    ser.marker.format.line.color.rgb = base
                except Exception:
                    pass
            elif kind == "area":
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = base
                ser.format.line.color.rgb = base
            elif is_pie:
                n = len(ser.points)
                for pi in range(n):
                    c = palette[pi % len(palette)] if len(palette) > 1 else base
                    ser.points[pi].format.fill.solid()
                    ser.points[pi].format.fill.fore_color.rgb = c
                    ser.points[pi].format.line.color.rgb = WHITE
                    ser.points[pi].format.line.width = Pt(1.2)
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = base
                ser.format.line.fill.background()
                if idxs:
                    for pi in range(len(ser.points)):
                        pt = ser.points[pi]
                        pt.format.fill.solid()
                        pt.format.fill.fore_color.rgb = (
                            highlight_color if pi in idxs else other_color)
                        pt.format.line.fill.background()
                        if show_values and pi in idxs:
                            try:
                                pt.data_label.font.size = Pt(value_size)
                                pt.data_label.font.name = F_MED
                                pt.data_label.font.bold = True
                                pt.data_label.font.color.rgb = PRIMARY
                            except Exception:
                                pass
        return ch

    def save(self, path):
        self.prs.save(path)
        print(f"saved: {path} | slides: {len(self.slides)}")
        return path


if __name__ == "__main__":
    if "--check" in sys.argv:
        t = find_base()
        print("base     :", t)
        print(f"size     : {os.path.getsize(t)/1024/1024:.1f} MB")
        prs = Presentation(t)
        print(f"canvas   : {prs.slide_width/914400:.3f} x {prs.slide_height/914400:.3f} in")
        print(f"masters  : {len(prs.slide_masters)} | layouts: "
              f"{[l.name for m in prs.slide_masters for l in m.slide_layouts]}")
        print(f"slides   : {len(prs.slides)}（将被清空）")
        print("icons    :", len(ensure_icons(t)), "个")
