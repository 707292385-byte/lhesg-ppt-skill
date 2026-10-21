# -*- coding: utf-8 -*-
"""pptx 版式渲染预览（替代 qlmanage）
按几何数据绘制每页布局示意：矩形/圆角/椭圆/自由形/图片/表格/组合 + 文字。
用法: python pptx_render_preview.py <pptx> <out_dir> [页码...]  (页码 1-based，省略=全部)
"""
import os, sys, math
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

DPI = 110
EMU_IN = 914400.0
HEI_M = "/System/Library/Fonts/Supplemental/STHeiti Medium.ttc"
HEI_L = "/System/Library/Fonts/Supplemental/STHeiti Light.ttc"
HEI_H = "/System/Library/Fonts/Supplemental/STHeiti Medium.ttc"
FONT_CACHE = {}


def font(path, px, weight=None):
    key = (path, px, weight)
    if key not in FONT_CACHE:
        try:
            FONT_CACHE[key] = ImageFont.truetype(path, px, index=0)
        except Exception:
            FONT_CACHE[key] = ImageFont.load_default()
    return FONT_CACHE[key]


def pick_font(name, bold):
    n = (name or "")
    if "Heavy" in n or "Bold" in n:
        return HEI_M
    if "Medium" in n:
        return HEI_M
    return HEI_L


def emu2px(v):
    return v / EMU_IN * DPI


def solid_fill(sh):
    """返回 fill 颜色 hex 或 None"""
    try:
        el = sh._element
        spPr = el.find(qn('p:spPr'))
        if spPr is None:
            spPr = el.find('.//' + qn('p:spPr'))
        if spPr is None:
            return None, None
        fill = spPr.find(qn('a:solidFill'))
        if fill is None:
            return None, None
        c = fill.find(qn('a:srgbClr'))
        if c is not None:
            return "#" + c.get('val'), None
        sch = fill.find(qn('a:schemeClr'))
        if sch is not None:
            m = {'accent1': '#5B9BD5', 'accent2': '#ED7D31', 'accent3': '#A5A5A5',
                 'accent4': '#FFC000', 'accent5': '#4472C4', 'accent6': '#70AD47',
                 'lt1': '#FFFFFF', 'dk1': '#000000', 'lt2': '#EEECE1', 'dk2': '#1F497D'}
            return m.get(sch.get('val'), '#CCCCCC'), None
    except Exception:
        pass
    # 渐变
    try:
        spPr = sh._element.find(qn('p:spPr'))
        grad = spPr.find(qn('a:gradFill'))
        if grad is not None:
            gs = grad.findall('.//' + qn('a:srgbClr'))
            if gs:
                return "#" + gs[0].get('val'), None
    except Exception:
        pass
    return None, None


def draw_shape_box(d, box, sh):
    x, y, w, h = box
    color, _ = solid_fill(sh)
    if color is None:
        return
    try:
        shp_type = sh.shape_type
    except Exception:
        shp_type = None
    st = str(shp_type)
    if "OVAL" in st or "椭圆" in (sh.name or ""):
        d.ellipse([x, y, x + w, y + h], fill=color)
    elif "ROUND" in st or "圆角" in (sh.name or ""):
        r = min(w, h) * 0.18
        d.rounded_rectangle([x, y, x + w, y + h], radius=r, fill=color)
    else:
        d.rectangle([x, y, x + w, y + h], fill=color)


def wrap(text, fnt, max_w, d):
    lines = []
    for raw in text.split("\n"):
        if not raw:
            lines.append("")
            continue
        cur = ""
        for ch in raw:
            t = cur + ch
            if d.textlength(t, font=fnt) > max_w and cur:
                lines.append(cur)
                cur = ch
            else:
                cur = t
        lines.append(cur)
    return lines


def render_text(d, box, sh):
    tf = sh.text_frame
    x, y, w, h = box
    pad = emu2px(45720)
    cx = x + pad
    avail = max(w - 2 * pad, 10)
    # 先收集各段行高，再按垂直锚点定位起始 y
    blocks = []
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        size = None
        fname = None
        bold = False
        col = "#333333"
        for r in p.runs:
            if r.font.size:
                size = r.font.size.pt
            if r.font.name:
                fname = r.font.name
            if r.font.bold:
                bold = True
            try:
                if r.font.color and r.font.color.rgb:
                    col = "#" + str(r.font.color.rgb)
            except Exception:
                pass
            break
        if size is None:
            size = 12
        fnt = font(pick_font(fname, bold), max(int(size * DPI / 72), 8))
        lines = wrap(txt, fnt, avail, d)
        lh = size * DPI / 72 * 1.28
        align = str(p.alignment) if p.alignment else "LEFT"
        blocks.append((lines, lh, align, fnt, col))
    total = sum(len(ls) * lh for ls, lh, _, _, _ in blocks)
    avail_h = max(h - 2 * pad, 10)
    anchor = tf.vertical_anchor
    if anchor is not None and "MIDDLE" in str(anchor):
        cy = y + pad + max((avail_h - total) / 2, 0)
    elif anchor is not None and "BOTTOM" in str(anchor):
        cy = y + h - pad - total
    else:
        cy = y + pad
    for lines, lh, align, fnt, col in blocks:
        for ln in lines:
            if not ln:
                cy += lh * 0.5
                continue
            if "CENTER" in align:
                tw = d.textlength(ln, font=fnt)
                d.text((cx + (avail - tw) / 2, cy), ln, font=fnt, fill=col)
            elif "RIGHT" in align:
                tw = d.textlength(ln, font=fnt)
                d.text((cx + avail - tw, cy), ln, font=fnt, fill=col)
            else:
                d.text((cx, cy), ln, font=fnt, fill=col)
            cy += lh


def walk(shapes, tf, out, d):
    for sh in shapes:
        try:
            st = sh.shape_type
        except Exception:
            st = None
        if st == MSO_SHAPE_TYPE.GROUP:
            grp = sh._element.find(qn('p:grpSpPr'))
            xfrm = grp.find(qn('a:xfrm')) if grp is not None else None
            if xfrm is None:
                continue
            off = xfrm.find(qn('a:off')); ext = xfrm.find(qn('a:ext'))
            co = xfrm.find(qn('a:chOff')); ce = xfrm.find(qn('a:chExt'))
            ox, oy = int(off.get('x')), int(off.get('y'))
            ex, ey = int(ext.get('cx')), int(ext.get('cy'))
            cox, coy = int(co.get('x')), int(co.get('y'))
            cex, cey = max(int(ce.get('cx')), 1), max(int(ce.get('cy')), 1)
            sx, sy = ex / cex, ey / cey

            def child_tf(X, Y, _tf=tf, ox=ox, oy=oy, cox=cox, coy=coy, sx=sx, sy=sy):
                px, py = _tf(ox + (X - cox) * sx, oy + (Y - coy) * sy)
                return px, py
            walk(sh.shapes, child_tf, out, d)
            continue

        try:
            l, t = int(sh.left), int(sh.top)
            w, h = int(sh.width), int(sh.height)
        except Exception:
            continue
        p1 = tf(l, t)
        p2 = tf(l + w, t + h)
        box = (p1[0], p1[1], p2[0] - p1[0], p2[1] - p1[1])

        if st == MSO_SHAPE_TYPE.PICTURE:
            try:
                img = sh.image
                import io
                im = Image.open(io.BytesIO(img.blob)).convert("RGBA")
                im = im.resize((max(int(box[2]), 1), max(int(box[3]), 1)))
                out.paste(im, (int(box[0]), int(box[1])), im)
            except Exception:
                d.rectangle([box[0], box[1], box[0] + box[2], box[1] + box[3]], outline="#BBBBBB", width=1)
            continue

        if sh.has_table:
            tbl = sh.table
            rows, cols = len(tbl.rows), len(tbl.columns)
            # 用真实行高/列宽，而不是平均切分
            col_ws = [c.width for c in tbl.columns]
            row_hs = [r.height for r in tbl.rows]
            tot_w, tot_h = sum(col_ws), sum(row_hs)
            if not tot_w or not tot_h:
                continue
            x0 = box[0]
            y0 = box[1]
            for i in range(rows):
                rh_px = box[3] * (row_hs[i] / tot_h)
                cx0 = box[0]
                for j in range(cols):
                    cw_px = box[2] * (col_ws[j] / tot_w)
                    c = tbl.cell(i, j)
                    bg = None
                    try:
                        bg = "#" + str(c.fill.fore_color.rgb)
                    except Exception:
                        bg = "#FFFFFF"
                    d.rectangle([cx0, y0, cx0 + cw_px, y0 + rh_px], fill=bg, outline="#DDDDDD")
                    # 取首个 run 的真实字号与颜色；按列宽换行，最多画到行高容纳的行数
                    sz_pt, col, bold, fname = 12.0, "#333333", False, None
                    paras = c.text_frame.paragraphs
                    for p in paras:
                        for r in p.runs:
                            if r.font.size:
                                sz_pt = r.font.size.pt
                            try:
                                if r.font.color and r.font.color.rgb:
                                    col = "#" + str(r.font.color.rgb)
                            except Exception:
                                pass
                            bold = bool(r.font.bold)
                            fname = r.font.name
                            break
                        break
                    fnt = font(pick_font(fname, bold), max(int(sz_pt * DPI / 72), 8), "b" if bold else None)
                    line_h = sz_pt * DPI / 72 * 1.3
                    max_lines = max(int(rh_px / line_h), 1)
                    # 逐段落换行（近似：按像素测宽），最多画行高容纳的行数
                    yy = y0 + 3
                    avail = cw_px - 10
                    stop = False
                    for p in paras:
                        if stop:
                            break
                        txt = "".join(r.text for r in p.runs)
                        if not txt:
                            continue
                        line = ""
                        for ch in txt:
                            if line and fnt.getlength(line + ch) > avail:
                                if max_lines <= 0:
                                    stop = True
                                    break
                                d.text((cx0 + 5, yy), line, font=fnt, fill=col)
                                yy += line_h
                                max_lines -= 1
                                line = ch
                            else:
                                line += ch
                        if line and not stop and max_lines > 0:
                            d.text((cx0 + 5, yy), line, font=fnt, fill=col)
                            yy += line_h
                            max_lines -= 1
                    cx0 += cw_px
                y0 += rh_px
            continue

        if getattr(sh, "has_chart", False):
            draw_chart(d, box, sh)
            continue

        if st in (MSO_SHAPE_TYPE.LINE, 9) or "LINE" in str(st):
            color, _ = solid_fill(sh)
            d.line([box[0], box[1] + box[3] / 2, box[0] + box[2], box[1] + box[3] / 2],
                   fill=color or "#999999", width=2)
            continue

        if getattr(sh, "has_text_frame", False) or st == MSO_SHAPE_TYPE.AUTO_SHAPE or st == MSO_SHAPE_TYPE.TEXT_BOX:
            draw_shape_box(d, box, sh)
            if sh.has_text_frame and sh.text_frame.text.strip():
                render_text(d, box, sh)


def _series_color(ser, i=None):
    """点色优先于系列色；取不到返回 None。"""
    if i is not None:
        try:
            c = ser.points[i].format.fill.fore_color.rgb
            if c:
                return "#" + str(c)
        except Exception:
            pass
    try:
        c = ser.format.fill.fore_color.rgb
        if c:
            return "#" + str(c)
    except Exception:
        pass
    try:
        c = ser.format.line.color.rgb
        if c:
            return "#" + str(c)
    except Exception:
        pass
    return None


def draw_chart(d, box, sh):
    """原生图表的近似绘制：条形/柱/折线/饼。供渲染自检用，不追求像素级还原。"""
    try:
        ch = sh.chart
        plot = ch.plots[0]
        cats = list(plot.categories)
        sers = list(ch.series)
        vals = [list(s.values) for s in sers]
        if not cats or not vals:
            raise ValueError
    except Exception:
        d.rectangle([box[0], box[1], box[0] + box[2], box[1] + box[3]],
                    outline="#CC3333", width=2)
        return
    x0, y0, W, H = box
    n = len(cats)
    show_lbl = plot.has_data_labels
    f_cat = font(HEI_L, max(int(11 * DPI / 72), 9))
    f_val = font(HEI_M, max(int(10 * DPI / 72), 9))
    GRAY = "#595959"

    root = sh.chart._chartSpace
    is_pie = root.find('.//' + qn('c:pieChart')) is not None or \
        root.find('.//' + qn('c:doughnutChart')) is not None
    line_el = root.find('.//' + qn('c:lineChart'))
    bar_el = root.find('.//' + qn('c:barChart'))
    horiz = bar_el is not None and bar_el.find(qn('c:barDir')) is not None \
        and bar_el.find(qn('c:barDir')).get('val') == 'bar'

    if is_pie:
        cx, cy = x0 + W / 2, y0 + H / 2
        r = min(W, H) / 2 * 0.92
        total = sum(vals[0]) or 1
        a0 = -90.0
        for i, v in enumerate(vals[0]):
            a1 = a0 + 360.0 * v / total
            col = _series_color(sers[0], i) or "#044AC0"
            d.pieslice([cx - r, cy - r, cx + r, cy + r], a0, a1, fill=col, outline="white", width=2)
            a0 = a1
        return

    if line_el is not None:
        mx = max((v for s in vals for v in s), default=1) or 1
        pad_l, pad_b, pad_t = 30, 22, 12
        px0, py0 = x0 + pad_l, y0 + pad_t
        pw, ph = W - pad_l - 14, H - pad_t - pad_b
        d.line([px0, py0 + ph, px0 + pw, py0 + ph], fill="#D6E2F5", width=1)
        d.line([px0, py0, px0, py0 + ph], fill="#D6E2F5", width=1)
        step = pw / max(n - 1, 1)
        for si, s in enumerate(sers):
            col = _series_color(s) or "#044AC0"
            pts = [(px0 + i * step, py0 + ph - (v / mx) * ph) for i, v in enumerate(vals[si])]
            d.line([p for pt in pts for p in pt], fill=col, width=3)
            for ptx, pty in pts:
                d.ellipse([ptx - 4, pty - 4, ptx + 4, pty + 4], fill=col)
        for i, c in enumerate(cats):
            d.text((px0 + i * step - f_cat.getlength(str(c)) / 2, py0 + ph + 4),
                   str(c), font=f_cat, fill=GRAY)
        if show_lbl:
            for i, v in enumerate(vals[0]):
                d.text((px0 + i * step - f_val.getlength(str(int(v))) / 2,
                        py0 + ph - (v / mx) * ph - 20), str(int(v)), font=f_val, fill=GRAY)
        return

    if horiz:
        # 类目标签预留宽度：按最长标签估算，封顶 32%
        lbl_w = max((f_cat.getlength(str(c)) for c in cats), default=40)
        pad_l = min(max(lbl_w + 10, 40), W * 0.32)
        pad_b = 8
        px0, py0 = x0 + pad_l, y0 + 4
        pw, ph = W - pad_l - 12, H - 8 - pad_b
        mx = max((v for s in vals for v in s), default=1) or 1
        d.line([px0, py0, px0, py0 + ph], fill="#D6E2F5", width=1)
        row_h = ph / n
        thickness = row_h * 0.62
        for i, c in enumerate(cats):
            cy = py0 + row_h * (i + 0.5)
            d.text((px0 - 8 - f_cat.getlength(str(c)), cy - f_cat.size / 2),
                   str(c), font=f_cat, fill=GRAY)
            for si, s in enumerate(sers):
                v = vals[si][i] or 0
                bw = (v / mx) * pw
                col = _series_color(s, i if len(sers) == 1 else None) or "#044AC0"
                d.rectangle([px0, cy - thickness / 2, px0 + bw, cy + thickness / 2], fill=col)
                if show_lbl:
                    d.text((px0 + bw + 6, cy - f_val.size / 2),
                           str(int(v)) if float(v).is_integer() else str(v),
                           font=f_val, fill="#595959")
        return

    # 纵向柱形
    pad_b = max(int(14 * DPI / 72), 22)
    pad_l, pad_t = 12, 16
    px0, py0 = x0 + pad_l, y0 + pad_t
    pw, ph = W - pad_l - 12, H - pad_t - pad_b
    mx = max((v for s in vals for v in s), default=1) or 1
    d.line([px0, py0 + ph, px0 + pw, py0 + ph], fill="#D6E2F5", width=1)
    col_w = pw / n
    thickness = col_w * 0.55
    for i, c in enumerate(cats):
        cx = px0 + col_w * (i + 0.5)
        d.text((cx - f_cat.getlength(str(c)) / 2, py0 + ph + 4), str(c), font=f_cat, fill=GRAY)
        for si, s in enumerate(sers):
            v = vals[si][i] or 0
            bh = (v / mx) * ph
            col = _series_color(s, i if len(sers) == 1 else None) or "#044AC0"
            d.rectangle([cx - thickness / 2, py0 + ph - bh, cx + thickness / 2, py0 + ph], fill=col)
            if show_lbl:
                d.text((cx - f_val.getlength(str(int(v))) / 2, py0 + ph - bh - 18),
                       str(int(v)), font=f_val, fill=GRAY)


def slide_bg(slide):
    """读取 p:bg：返回 ("blip", bytes) 或 ("solid", "RRGGBB") 或 None。"""
    cSld = slide._element.find(qn("p:cSld"))
    if cSld is None:
        return None
    bg = cSld.find(qn("p:bg"))
    if bg is None:
        return None
    blip = bg.find(".//" + qn("a:blip"))
    if blip is not None:
        rId = blip.get(qn("r:embed"))
        try:
            return ("blip", slide.part.related_part(rId).blob)
        except Exception:
            return None
    clr = bg.find(".//" + qn("a:srgbClr"))
    return ("solid", clr.get("val")) if clr is not None else None


def render(pptx, idx, out_png):
    prs = Presentation(pptx)
    W = int(emu2px(prs.slide_width))
    H = int(emu2px(prs.slide_height))
    slide = prs.slides[idx]
    out = Image.new("RGB", (W, H), "white")
    bg = slide_bg(slide)
    if bg:
        kind, val = bg
        if kind == "solid":
            out = Image.new("RGB", (W, H), "#" + val)
        else:
            import io
            pic = Image.open(io.BytesIO(val)).convert("RGB")
            out.paste(pic.resize((W, H), Image.LANCZOS), (0, 0))
    d = ImageDraw.Draw(out)

    def ident(x, y):
        return emu2px(x), emu2px(y)

    walk(slide.shapes, ident, out, d)
    out.save(out_png)
    return out_png


if __name__ == "__main__":
    path = sys.argv[1]
    outdir = sys.argv[2]
    os.makedirs(outdir, exist_ok=True)
    prs = Presentation(path)
    pages = [int(a) for a in sys.argv[3:]] or list(range(1, len(prs.slides) + 1))
    for p in pages:
        f = os.path.join(outdir, f"slide{p:02d}.png")
        render(path, p - 1, f)
        print("rendered", f)
