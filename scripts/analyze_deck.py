# -*- coding: utf-8 -*-
"""
analyze_deck.py — 从"被认可的 deck"反解设计经验（把成品变成可迁移参数）

用途：当你认可某一版 deck 的排版布局时，用本脚本提取其字号、配色、几何锚点、
     版式节奏等统计特征，据此更新 references/design-experience.md。
     固化的对象是这些**规律**，不是那个文件本身。

用法: python analyze_deck.py <deck.pptx>
"""
import sys
from collections import Counter

from pptx import Presentation
from pptx.oxml.ns import qn


def inch(v):
    return round(v / 914400, 2) if v is not None else None


def main(path):
    prs = Presentation(path)
    sizes, fills, fonts, text_colors = Counter(), Counter(), Counter(), Counter()
    xs, widths, heights, ys, radius = (Counter() for _ in range(5))

    print(f"deck : {path}")
    print(f"页数 : {len(prs.slides)} | 画布 {prs.slide_width/914400:.3f}"
          f"×{prs.slide_height/914400:.3f} in\n")

    for pi, slide in enumerate(prs.slides, 1):
        n_sh = n_pic = n_tbl = n_txt = 0
        tsize = []
        for sh in slide.shapes:
            n_sh += 1
            st = str(sh.shape_type)
            if "PICTURE" in st:
                n_pic += 1
            if sh.has_table:
                n_tbl += 1
                for r in sh.table.rows:
                    for c in r.cells:
                        for p in c.text_frame.paragraphs:
                            for run in p.runs:
                                if run.font.size:
                                    sizes[round(run.font.size.pt, 1)] += 1
                                    tsize.append(run.font.size.pt)
            if sh.has_text_frame:
                n_txt += 1
                for p in sh.text_frame.paragraphs:
                    for run in p.runs:
                        if run.font.size:
                            sizes[round(run.font.size.pt, 1)] += 1
                            tsize.append(run.font.size.pt)
                        if run.font.name:
                            fonts[run.font.name] += 1
                        try:
                            if run.font.color and run.font.color.rgb:
                                text_colors[str(run.font.color.rgb)] += 1
                        except Exception:
                            pass
            try:
                spPr = sh._element.find(qn('p:spPr'))
                if spPr is not None:
                    f = spPr.find(qn('a:solidFill'))
                    if f is not None:
                        c = f.find(qn('a:srgbClr'))
                        if c is not None:
                            fills[c.get('val')] += 1
                    prst = spPr.find(qn('a:prstGeom'))
                    if prst is not None and prst.get('prst') == 'roundRect':
                        av = prst.find(qn('a:avLst'))
                        if av is not None:
                            gd = av.find(qn('a:gd'))
                            if gd is not None:
                                radius[gd.get('fmla').replace('val ', '')] += 1
            except Exception:
                pass
            l, t, w, h = inch(sh.left), inch(sh.top), inch(sh.width), inch(sh.height)
            if l is not None:
                xs[l] += 1
            if w is not None:
                widths[w] += 1
            if h is not None:
                heights[h] += 1
            if t is not None:
                ys[t] += 1
        tsz = f"字号 {min(tsize):.1f}-{max(tsize):.1f}" if tsize else "无文字"
        print(f"P{pi:<2} 形状{n_sh:<3} 图{n_pic} 表{n_tbl} 文本框{n_txt:<3} {tsz}")

    print("\n=== 字号频次 ===")
    for s, c in sorted(sizes.items()):
        print(f"  {s:>5}  {c}")
    body_like = [s for s in sizes if 9 <= s <= 14]
    if body_like:
        mid = sorted(body_like)[len(body_like)//2]
        if mid < 12:
            print(f"  ⚠️ 正文中位约 {mid}pt，低于规范下限 12pt → 设计时按 14pt（密集页 12pt）执行")

    print("\n=== 字体 ===", dict(fonts))
    print("\n=== 字色 top ===",
          [(f"#{c}", n) for c, n in text_colors.most_common(10)])
    print("\n=== 填充色 top ===",
          [(f"#{c}", n) for c, n in fills.most_common(12)])
    print("\n=== 几何锚点 ===")
    print("  左边界 x :", xs.most_common(8))
    print("  宽度 w   :", widths.most_common(8))
    print("  高度 h   :", heights.most_common(10))
    print("  顶边 y   :", ys.most_common(8))
    print("  圆角档   :", radius.most_common(6))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    main(sys.argv[1])
